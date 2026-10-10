"""Format natif, écritures atomiques et échanges Word simplifiés."""
import base64
import io
import json
import os
import tempfile
from html import escape
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph
from docx.table import Table
from PySide6.QtCore import Qt, QUrl
from PySide6.QtGui import QTextDocument, QTextTable, QFont


def atomic_write(path, text):
    path = Path(path)
    fd, tmp = tempfile.mkstemp(dir=path.parent, prefix='.plume-', suffix='.tmp')
    try:
        with os.fdopen(fd, 'w', encoding='utf-8') as f:
            f.write(text)
            f.flush()
            os.fsync(f.fileno())
        os.replace(tmp, path)
    finally:
        Path(tmp).unlink(missing_ok=True)


def validate_project(data):
    if not isinstance(data, dict) or data.get('version') != 1:
        raise ValueError('Ce fichier n’est pas un document Plume compatible.')
    chapters = data.get('chapters')
    if not isinstance(chapters, list) or not chapters or len(chapters) > 10000:
        raise ValueError('La liste des chapitres est invalide.')
    for chapter in chapters:
        if not isinstance(chapter, dict) or not all(isinstance(chapter.get(k), str) for k in ('title', 'html')):
            raise ValueError('Un chapitre est endommagé.')
    return data


def save_project(path, data):
    validate_project(data)
    atomic_write(path, json.dumps(data, ensure_ascii=False, indent=2))


def load_project(path):
    return validate_project(json.loads(Path(path).read_text(encoding='utf-8')))


def import_docx(path):
    doc = Document(path)
    def para(p):
        chunks = []
        for run in p.runs:
            css = []
            if run.bold: css.append('font-weight:700')
            if run.italic: css.append('font-style:italic')
            if run.underline: css.append('text-decoration:underline')
            if run.font.size: css.append(f'font-size:{run.font.size.pt}pt')
            if run.font.name: css.append('font-family:' + escape(run.font.name, quote=True))
            if run.font.color.rgb: css.append('color:#' + str(run.font.color.rgb))
            chunks.append(f'<span style="{";".join(css)}">{escape(run.text).replace(chr(10), "<br>")}</span>')
            for blip in run._r.xpath('.//a:blip'):
                rid = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                if rid and rid in doc.part.related_parts:
                    part = doc.part.related_parts[rid]
                    uri = f'data:{part.content_type};base64,{base64.b64encode(part.blob).decode()}'
                    chunks.append(f'<img src="{uri}" width="420">')
        level = 0
        name = p.style.name if p.style else ''
        if name.startswith('Heading ') and name[-1:].isdigit(): level = min(6, int(name[-1]))
        tag = f'h{level}' if level else 'p'
        align = {WD_ALIGN_PARAGRAPH.CENTER:'center', WD_ALIGN_PARAGRAPH.RIGHT:'right', WD_ALIGN_PARAGRAPH.JUSTIFY:'justify'}.get(p.alignment, 'left')
        return f'<{tag} style="text-align:{align}">{"".join(chunks) or "<br>"}</{tag}>'
    parts = []
    for node in doc.element.body.iterchildren():
        if node.tag.endswith('}p'): parts.append(para(Paragraph(node, doc)))
        elif node.tag.endswith('}tbl'):
            table = Table(node, doc)
            rows = ['<tr>' + ''.join('<td>' + ''.join(para(p) for p in cell.paragraphs) + '</td>' for cell in row.cells) + '</tr>' for row in table.rows]
            parts.append('<table border="1" cellspacing="0" cellpadding="6">' + ''.join(rows) + '</table>')
    return '<html><body>' + ''.join(parts) + '</body></html>'


def export_docx(html, path):
    """Export paragraphs, run styling, lists, images and simple tables."""
    qt = QTextDocument()
    qt.setHtml(html)
    out = Document()
    section = out.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin = section.right_margin = Mm(20)

    def paragraph(block, parent):
        fmt = block.blockFormat()
        level = fmt.headingLevel()
        style = f'Heading {min(level, 9)}' if level else None
        if block.textList():
            from PySide6.QtGui import QTextListFormat
            style = 'List Number' if block.textList().format().style() in (QTextListFormat.ListDecimal, QTextListFormat.ListLowerAlpha, QTextListFormat.ListUpperAlpha) else 'List Bullet'
        p = parent.add_paragraph(style=style)
        a = fmt.alignment()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if a & Qt.AlignJustify else WD_ALIGN_PARAGRAPH.CENTER if a & Qt.AlignHCenter else WD_ALIGN_PARAGRAPH.RIGHT if a & Qt.AlignRight else WD_ALIGN_PARAGRAPH.LEFT
        it = block.begin()
        while not it.atEnd():
            frag = it.fragment()
            if frag.isValid():
                cf = frag.charFormat()
                if cf.isImageFormat():
                    name = cf.toImageFormat().name()
                    try:
                        if name.startswith('data:'):
                            blob = base64.b64decode(name.split(',', 1)[1])
                        else:
                            image = qt.resource(QTextDocument.ImageResource, QUrl(name))
                            from PySide6.QtCore import QBuffer, QByteArray, QIODevice
                            data = QByteArray(); buffer = QBuffer(data); buffer.open(QIODevice.WriteOnly)
                            if image is None or not image.save(buffer, 'PNG'): raise ValueError('Image illisible')
                            blob = bytes(data)
                        width = cf.toImageFormat().width() or 420
                        p.add_run().add_picture(io.BytesIO(blob), width=Mm(min(170, width * 25.4 / 96)))
                    except Exception as exc:
                        raise ValueError('Une image ne peut pas être exportée en Word.') from exc
                else:
                    r = p.add_run(frag.text().replace('\u2028', '\n'))
                    r.bold = cf.fontWeight() >= QFont.Bold
                    r.italic = cf.fontItalic(); r.underline = cf.fontUnderline()
                    r.font.strike = cf.fontStrikeOut()
                    r.font.name = cf.font().family() or 'Calibri'
                    if cf.fontPointSize(): r.font.size = Pt(cf.fontPointSize())
                    color = cf.foreground().color()
                    if cf.foreground().style() != Qt.NoBrush:
                        r.font.color.rgb = RGBColor(color.red(), color.green(), color.blue())
            it += 1
        return p

    def walk(frame, parent):
        it = frame.begin()
        while not it.atEnd():
            child = it.currentFrame()
            if isinstance(child, QTextTable):
                table = parent.add_table(rows=child.rows(), cols=child.columns())
                table.style = 'Table Grid'
                for row in range(child.rows()):
                    for col in range(child.columns()):
                        cell = child.cellAt(row, col)
                        if cell.rowSpan() != 1 or cell.columnSpan() != 1:
                            raise ValueError('Les cellules fusionnées ne sont pas prises en charge dans cet export Word.')
                        target = table.cell(row, col)
                        ci = cell.begin()
                        while not ci.atEnd():
                            if ci.currentBlock().isValid(): paragraph(ci.currentBlock(), target)
                            ci += 1
                        if len(target.paragraphs) > 1:
                            el = target.paragraphs[0]._element; el.getparent().remove(el)
            elif child: walk(child, parent)
            elif it.currentBlock().isValid(): paragraph(it.currentBlock(), parent)
            it += 1
    walk(qt.rootFrame(), out)
    out.save(path)
