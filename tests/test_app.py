import os
os.environ['QT_QPA_PLATFORM'] = 'offscreen'
import sys
import tempfile
import unittest
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from PySide6.QtWidgets import QApplication
from PySide6.QtGui import QImage, QTextCursor, QTextDocument
from PySide6.QtCore import QBuffer, QByteArray, QIODevice, QStandardPaths
from docx import Document
from storage import save_project, load_project, export_docx, import_docx
from main import Window
app = QApplication.instance() or QApplication([])
app.setApplicationName('Plume-tests')
QStandardPaths.setTestModeEnabled(True)

class Tests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.w = Window()
        self.w.timer.stop()
        self.w.recovery = self.root / 'recovery.plume'
    def tearDown(self):
        self.w.dirty = False
        self.w.close()
        self.tmp.cleanup()
    def test_native_chapters_images_recovery(self):
        self.w.editor.setHtml('<h1>Bonjour</h1><p><b>Été</b> et histoires.</p>')
        image = QImage(20, 10, QImage.Format_RGB32); image.fill(0xffaa44)
        data = QByteArray(); buf = QBuffer(data); buf.open(QIODevice.WriteOnly); image.save(buf, 'PNG')
        cursor = self.w.editor.textCursor(); cursor.movePosition(QTextCursor.End); self.w.editor.setTextCursor(cursor)
        self.w.editor.embed(bytes(data), 100)
        self.w.chapters.append({'title':'Deux', 'html':'<p>Seconde partie.</p>'})
        self.w.switch_chapter(1); self.w.editor.insertPlainText('Test ')
        self.w.path = self.root / 'book.plume'
        self.assertTrue(self.w.save())
        saved = load_project(self.w.path)
        self.assertEqual(len(saved['chapters']), 2)
        self.assertIn('data:image/png;base64,', saved['chapters'][0]['html'])
        doc = QTextDocument(); doc.setHtml(saved['chapters'][0]['html'])
        self.assertIn('Été', doc.toPlainText())
        self.w.editor.insertPlainText('non enregistré')
        self.w.autosave()
        self.assertIn('non enregistré', load_project(self.w.recovery)['chapters'][1]['html'])
    def test_pdf_and_word(self):
        self.w.editor.setHtml('<h1>Une histoire</h1><p><b>Texte gras</b> et <i>italique</i>.</p><table border="1"><tr><td>Cellule A</td><td>Cellule B</td></tr></table>')
        pdf = self.root/'test.pdf'; self.w.combined_document().print_(self.w.printer(pdf))
        self.assertTrue(pdf.read_bytes().startswith(b'%PDF'))
        word = self.root/'test.docx'; export_docx(self.w.combined_document().toHtml(), word)
        d = Document(word)
        self.assertEqual(len(d.tables), 1)
        self.assertEqual(d.tables[0].cell(0,0).text, 'Cellule A')
        self.assertTrue(any(r.bold and 'Texte gras' in r.text for p in d.paragraphs for r in p.runs))
        imported = import_docx(word); self.assertIn('Cellule B', imported); self.assertIn('font-weight:700', imported)
    def test_invalid_file_does_not_overwrite(self):
        path = self.root/'invalid.plume'; path.write_text('original')
        with self.assertRaises(ValueError): save_project(path, {'version':99})
        self.assertEqual(path.read_text(), 'original')
    def test_embedded_image_word_roundtrip(self):
        image = QImage(20, 10, QImage.Format_RGB32); image.fill(0xffaa44)
        data = QByteArray(); buf = QBuffer(data); buf.open(QIODevice.WriteOnly); image.save(buf, 'PNG')
        self.w.editor.embed(bytes(data), 100)
        word = self.root/'image.docx'; export_docx(self.w.editor.toHtml(), word)
        self.assertEqual(len(Document(word).inline_shapes), 1)
        imported = import_docx(word); self.assertIn('data:image/png;base64,', imported)
    def test_screenshot(self):
        self.w.editor.setHtml('<h1>Chaque histoire commence ici.</h1><p>Bienvenue dans <b>Plume</b>, votre espace pour écrire, imaginer et créer.</p><h2>Un endroit pour vos idées</h2><p>Une lettre, un chapitre, un personnage… Prenez le temps de trouver vos mots.</p>')
        self.w.show(); app.processEvents()
        self.w.grab().save(str(Path(__file__).parents[1]/'apercu.png'))

if __name__ == '__main__': unittest.main()
